#!/usr/bin/env python3
"""
Zero to Pro — Professional Project Presentation Generator.

Builds ztp.docx (a Word-format slide deck) summarizing the project:
architecture, features, tech stack, compiler engine, quiz system,
AI assistant, admin panel, security, QA/test results and roadmap.

Dependency: python-docx  (pip install python-docx)
Run:         python scripts/generate-presentation.py
"""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Palette (matches the app's indigo theme) ────────────────────────────────
INDIGO = RGBColor(0x43, 0x38, 0xca)
ACCENT = RGBColor(0x6a, 0x5c, 0xff)
DARK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x6b, 0x72, 0x80)
WHITE = RGBColor(0xff, 0xff, 0xff)
GREEN = RGBColor(0x10, 0xb9, 0x81)
RED = RGBColor(0xef, 0x44, 0x44)


def shade_cell(cell, hex_color):
    """Apply a background fill to a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def run_font(run, size=11, bold=False, italic=False, color=DARK):
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"


def add_title(doc, text, size=30, color=INDIGO, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    run_font(r, size=size, bold=True, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_heading_block(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    run_font(r, size=17, bold=True, color=ACCENT)
    if subtitle:
        sp = doc.add_paragraph()
        s = sp.add_run(subtitle)
        run_font(s, size=10.5, italic=True, color=GRAY)
        sp.paragraph_format.space_after = Pt(6)
    return p


def add_para(doc, text, size=11, color=DARK, bold=False, italic=False, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    run_font(r, size=size, bold=bold, italic=italic, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items, size=11):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            label, rest = it
            r = p.add_run(label)
            run_font(r, size=size, bold=True, color=DARK)
            r2 = p.add_run(rest)
            run_font(r2, size=size, color=DARK)
        else:
            r = p.add_run(it)
            run_font(r, size=size, color=DARK)
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        pr = hdr[i].paragraphs[0]
        run = pr.add_run(h)
        run_font(run, size=10.5, bold=True, color=WHITE)
        pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shade_cell(hdr[i], "4338CA")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            pr = cells[i].paragraphs[0]
            run = pr.add_run(str(val))
            run_font(run, size=10)
            if i == len(row) - 1 and isinstance(val, str) and ("✓" in val or val == "Fixed"):
                run_font(run, size=10, bold=True, color=GREEN)
    if col_widths:
        for i, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


def new_section(doc):
    doc.add_page_break()


def build_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Zero to Pro · Zero-to-Pro Platform · Confidential")
    run_font(r, size=8, color=GRAY)


def main():
    doc = Document()
    # Base style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # ── Slide 1 · Title ─────────────────────────────────────────────────────
    for _ in range(3):
        doc.add_paragraph()
    add_title(doc, "ZERO TO PRO", size=40, color=INDIGO)
    add_title(doc, "Full-Stack Coding Learning Platform", size=18, color=GRAY)
    doc.add_paragraph()
    add_title(doc, "AI Tutor  ·  Live Compiler  ·  Quizzes & Certificates", size=13, color=ACCENT)
    doc.add_paragraph()
    add_para(doc, "Project Presentation  ·  v2.0", size=10, color=GRAY, italic=True,
             space_after=0)
    add_para(doc, "September 2026", size=10, color=GRAY, italic=True)
    new_section(doc)

    # ── Slide 2 · Project Overview ──────────────────────────────────────────
    add_heading_block(doc, "1 · Project Overview",
                      "What is Zero to Pro and why it was built")
    add_para(doc,
             "Zero to Pro is a complete learning platform that takes a beginner from "
             "zero coding knowledge to a confident programmer. Instead of separate tools, "
             "everything a student needs lives in one web application: guided tutorials, "
             "language quizzes, a real live code compiler, verifiable certificates and an "
             "always-available AI coding tutor.")
    add_bullets(doc, [
        ("End-to-end platform — ", "students can read, practice, test and get certified without leaving the app."),
        ("Real execution — ", "the compiler runs genuine Python, Java, JavaScript and PHP programs server-side."),
        ("AI-assisted learning — ", "an embedded assistant explains concepts and auto-fixes compiler errors."),
        ("Gamified progress — ", "passing a quiz in any language earns a branded certificate."),
    ])
    new_section(doc)

    # ── Slide 3 · Key Features ──────────────────────────────────────────────
    add_heading_block(doc, "2 · Key Features",
                      "The modules that make up the product")
    add_bullets(doc, [
        ("📚 Interactive Tutorials — ", "70+ step-by-step lessons across 7 languages, each code snippet paired with a plain-language explanation."),
        ("🎯 Quiz Engine — ", "800 questions across 8 categories; 35 random questions per attempt; 75% required to pass."),
        ("🏆 Certificate System — ", "automatically issues (and updates on retake) a branded certificate per language."),
        ("⚙️ Live Compiler — ", "Monaco Editor with output console, stdin input and live browser previews for HTML/CSS/React/Bootstrap."),
        ("🤖 AI Assistant — ", "Gemini-powered widget, server-side key, learns each student's language and tone, auto-opens on errors."),
        ("👑 Admin Panel — ", "manage users, toggle roles and create/update/delete tutorials."),
        ("🌗 Theme Switch — ", "light/dark theme, persisted per user."),
    ])
    new_section(doc)

    # ── Slide 4 · Technology Stack ──────────────────────────────────────────
    add_heading_block(doc, "3 · Technology Stack",
                      "The tools and runtimes behind the platform")
    add_table(
        doc,
        ["Layer", "Technology", "Purpose"],
        [
            ["Backend", "Node.js 18+ · Express 5", "HTTP server, routing, middleware"],
            ["Database", "MongoDB · Mongoose 9", "Users, tests, tutorials persistence"],
            ["Templating", "EJS · Tailwind CSS", "Server-rendered pages & styling"],
            ["Editor", "Monaco Editor (CDN)", "In-browser code editing"],
            ["Auth", "express-session · bcryptjs", "Session-based login, hashed passwords"],
            ["Compiler", "Node · Python · Java · PHP (XAMPP)", "Real code execution + previews"],
            ["AI", "Google Gemini (Flash)", "Coding tutor and error fixing"],
        ],
        col_widths=[1.1, 2.2, 2.7],
    )
    new_section(doc)

    # ── Slide 5 · Architecture ──────────────────────────────────────────────
    add_heading_block(doc, "4 · System Architecture",
                      "Clean Model–View–Controller layout")
    add_para(doc, "The codebase follows the MVC pattern, keeping data, logic and presentation "
                  "cleanly separated:")
    add_table(
        doc,
        ["Tier", "Folder", "Responsibility"],
        [
            ["Models", "models/", "Mongoose schemas — User, Test, Home (tutorials)"],
            ["Views", "views/", "EJS templates — pages, layouts and partials"],
            ["Controllers", "controllers/", "Request handlers, business logic, validation"],
            ["Routes", "routes/", "URL mapping with auth/role guards"],
            ["Static", "public/ · logo/", "Client-side JS, CSS and brand assets"],
        ],
        col_widths=[1.0, 1.9, 3.1],
    )
    add_para(doc, "Request flow:", size=11, bold=True, space_after=2)
    add_para(doc, "Browser → Express route (guard check) → Controller → Mongoose/Model → "
                  "EJS view → HTML response", size=10.5, italic=True, color=GRAY)
    new_section(doc)

    # ── Slide 6 · Compiler Engine ───────────────────────────────────────────
    add_heading_block(doc, "5 · Live Compiler Engine",
                      "Nine languages, two execution models")
    add_table(
        doc,
        ["Language", "Execution", "Notes"],
        [
            ["Python 3", "Server-side", "Full standard-logic support, stdin input"],
            ["Java (JDK 21)", "Server-side", "Compiles + runs named public class"],
            ["JavaScript", "Server-side", "Node.js runtime, safe built-ins only"],
            ["PHP 8 (XAMPP)", "Server-side", "Auto-detected install"],
            ["C++ (g++/C++17)", "Server-side", "Gives a friendly guide if g++ is missing"],
            ["HTML / CSS", "Browser preview", "Live iframe rendering"],
            ["React (JSX)", "Browser preview", "Babel-transpiled in the iframe"],
            ["Bootstrap", "Browser preview", "CDN + live preview"],
        ],
        col_widths=[1.5, 1.7, 2.8],
    )
    add_para(doc, "Hardcore stress code (fibonacci, prime sieve, sorting, recursion, "
                  "feedback loops) was executed successfully during QA — see Testing section.",
             size=10, italic=True, color=GRAY)
    new_section(doc)

    # ── Slide 7 · Compiler Security ─────────────────────────────────────────
    add_heading_block(doc, "6 · Compiler Security Model",
                      "Safe execution for an open classroom")
    add_bullets(doc, [
        ("Module allow-lists — ", "Python and Node only accept a whitelist of safe standard modules (fs / child_process / network are blocked)."),
        ("Command blocking — ", "PHP forbids system / exec / shell_exec / passthru / popen / proc_open."),
        ("Java import rule — ", "only java.* / javax.* imports are permitted."),
        ("Resource limits — ", "30-second timeout and a 20 MB output cap prevent runaway programs."),
        ("Session protection — ", "the compiler API is only reachable by logged-in users."),
    ])
    add_bullets(doc, [("Verified live:", " forbidden imports and system() calls were rejected with clear security-error messages.")])
    new_section(doc)

    # ── Slide 8 · Quiz & Certificates ───────────────────────────────────────
    add_heading_block(doc, "7 · Quiz & Certification System",
                      "Assessment with real recognition")
    add_table(
        doc,
        ["Category", "Questions", "qNum Range"],
        [
            ["HTML", "100", "1 – 100"],
            ["CSS", "100", "101 – 200"],
            ["JavaScript", "100", "201 – 300"],
            ["PHP", "100", "301 – 400"],
            ["C++", "100", "401 – 500"],
            ["React", "100", "501 – 600"],
            ["Bootstrap", "100", "601 – 700"],
            ["General", "100", "701 – 800"],
        ],
        col_widths=[1.8, 2.0, 2.2],
    )
    add_para(doc, "Per attempt: 35 random questions · passing score 75% · certificates are "
                  "keyed by language and updated (not duplicated) on a successful retake.",
             size=10.5, italic=True, color=GRAY)
    new_section(doc)

    # ── Slide 9 · AI Assistant ──────────────────────────────────────────────
    add_heading_block(doc, "8 · AI Coding Assistant",
                      "Gemini-powered tutoring, server-side")
    add_bullets(doc, [
        ("Identity-safe persona — ", "introduces itself as “Zero to Pro AI, built by the ZTP team”; never misattributes itself to other vendors."),
        ("Bilingual replies — ", "answers in simple English or Roman Urdu/Hindi to match the student."),
        ("Error triage — ", "when a compiled program fails, the AI opens by itself, explains the fault and rewrites the corrected code."),
        ("Server-side key — ", "the Gemini API key never reaches the browser; the widget calls a protected /api/ai/chat endpoint."),
        ("Retry resilience — ", "transient Gemini 503/429 responses are retried with backoff."),
    ])
    new_section(doc)

    # ── Slide 10 · Admin Panel ──────────────────────────────────────────────
    add_heading_block(doc, "9 · Admin Panel",
                      "Safe management of the community")
    add_bullets(doc, [
        ("User management — ", "view all accounts, toggle user ↔ admin roles, delete users."),
        ("Self-protection — ", "admins cannot demote or delete their own account."),
        ("Last-admin guard — ", "the final remaining admin can never be deleted."),
        ("Tutorial CRUD — ", "add, edit and delete tutorials with structured, sectioned content."),
        ("Flash feedback — ", "every action gives a clear success/error message."),
    ])
    new_section(doc)

    # ── Slide 11 · Database Schema ──────────────────────────────────────────
    add_heading_block(doc, "10 · Database Schema",
                      "Three core collections")
    add_table(
        doc,
        ["Collection", "Key Fields", "Notes"],
        [
            ["users", "name, email, password (bcrypt), role", "role: user | admin"],
            ["users.certificates[]", "category, testTitle, percentage, certIdZZH-XXX-XXXXXX, earnedAt", "one per language"],
            ["tests", "title, questions[]", "800 questions with fixed qNum buckets"],
            ["questions[]", "qNum, questionText, options[4], correctAnswerIndex", "category tag for display"],
            ["homes", "title, description, category, content, sections[]", "structured tutorials"],
        ],
        col_widths=[1.7, 3.1, 1.2],
    )
    new_section(doc)

    # ── Slide 12 · Security Hardening ───────────────────────────────────────
    add_heading_block(doc, "11 · Security Hardening",
                      "Protecting users and data")
    add_bullets(doc, [
        ("Password hashing — ", "bcrypt with salt before any user document is written."),
        ("Door-guarded routes — ", "requiresAuth / requireAdmin middleware on every protected path."),
        ("Input validation — ", "signup validates name, Gmail format, password length and confirmation; quiz/API payloads are sanity-checked."),
        ("Clean HTTP semantics — ", "404/400/500 responses for both HTML and JSON API consumers."),
        ("Secrets hygiene — ", ".env is git-ignored; a documented .env.example replaces committed secrets."),
        ("Compile-time abuse control — ", "module allow-lists and command blocking inside the compiler (slide 6)."),
    ])
    new_section(doc)

    # ── Slide 13 · Testing & QA ─────────────────────────────────────────────
    add_heading_block(doc, "12 · Testing & Quality Assurance",
                      "Every area was exercised on a live server")
    add_table(
        doc,
        ["Test case", "Input", "Result"],
        [
            ["Python — fibonacci + prime sieve", "hardcore 50-number sieve", "PASS ✓"],
            ["Java — bubble sort", "array of 7 ints", "PASS ✓"],
            ["JavaScript — recursive factorial + closures", "10!", "PASS ✓"],
            ["PHP — prime loop + JSON", "primes 1–100", "PASS ✓"],
            ["C++ — missing g++", "hello world", "Friendly guide ✓"],
            ["stdin piping", "Python a+b", "PASS ✓"],
            ["Python forbidden import", "import requests", "Blocked ✓"],
            ["PHP system() call", "system(…)", "Blocked ✓"],
            ["AI identity question", "“Who are you?”", "PASS ✓"],
            ["Invalid test id", "POST /api/test/someid", "404 ✓"],
        ],
        col_widths=[2.4, 2.0, 1.6],
    )
    add_para(doc, "Server start, authentication flows, quiz submission, certificates pages, "
                  "admin dashboard and the 404/error handlers were also verified end-to-end.",
             size=10, italic=True, color=GRAY)
    new_section(doc)

    # ── Slide 14 · Bugs Found & Fixed ───────────────────────────────────────
    add_heading_block(doc, "13 · Bugs Found & Fixed",
                      "Identified during the project audit")
    add_table(
        doc,
        ["Issue", "Impact", "Resolution"],
        [
            ["General category (701–800) missing from quiz ranges", "100 questions unreachable", "Fixed — range + certificate code added"],
            ["Invalid ObjectId → 500 CastError", "confusing server error", "Fixed — clean 404 response"],
            ["No password-length rule on signup", "weak accounts", "Fixed — min 6 chars validation"],
            ["Unused deps (axios, express-validator)", "bloat", "Removed from package.json"],
            ["Missing .env.example", "poor onboarding", "Added with documented vars"],
            ["/api 404 returned HTML", "inconsistent API", "Fixed — JSON error for APIs"],
        ],
        col_widths=[2.6, 1.5, 1.9],
    )
    new_section(doc)

    # ── Slide 15 · Roadmap ──────────────────────────────────────────────────
    add_heading_block(doc, "14 · Future Scope",
                      "Where the platform goes next")
    add_bullets(doc, [
        ("Online compile API — ", "hook Piston/Judge0 so C++ (and more languages) run without a local install."),
        ("CodePlayground leaderboards — ", "scores, rankings and streaks for the quiz engine."),
        ("PDF certificates — ", "downloadable branded PDFs via pdfkit."),
        ("Progress analytics — ", "per-student learning dashboards for admins."),
        ("Rate limiting — ", "protect login and AI endpoints from abuse."),
        ("Deployment — ", "Docker compose + CI pipeline for reproducible hosting."),
    ])
    new_section(doc)

    # ── Slide 16 · Closing ──────────────────────────────────────────────────
    add_title(doc, "Thank You", size=34, color=INDIGO)
    doc.add_paragraph()
    add_title(doc, "Zero to Pro — Learn, Practice, Get Certified.", size=15, color=DARK)
    doc.add_paragraph()
    add_para(doc, "Built by the Zero to Pro team", size=11, color=GRAY, italic=True,
             space_after=0)
    add_para(doc, "Questions & feedback welcome.", size=10, color=GRAY, italic=True)

    build_footer(doc)

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "ztp.docx")
    out = os.path.normpath(out)
    doc.save(out)
    print("Saved:", out)


if __name__ == "__main__":
    main()